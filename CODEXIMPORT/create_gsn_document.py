from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).with_name("GSN-project_draadloze_3D-meetmodule.docx")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell, color="8EA9C1"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:color"), color)


def add_box(cell, label, fill="EAF2F8"):
    cell.text = ""
    shade(cell, fill)
    borders(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(label)
    run.bold = True
    run.font.size = Pt(9)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.65)
section.bottom_margin = Cm(1.55)
section.left_margin = Cm(1.75)
section.right_margin = Cm(1.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(5)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("GSN-project: draadloze 3D-meetmodule\nmet vliegtuig-mock-up")
run.bold = True
run.font.name = "Aptos Display"
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(31, 78, 121)
title.paragraph_format.space_after = Pt(11)

def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12.5)
    r.font.color.rgb = RGBColor(31, 78, 121)


heading("Opdrachtomschrijving")
paragraphs = [
    "Voor mijn GSN-project ontwerp en bouw ik een compacte, draadloze meetmodule in een zelf ontworpen en 3D-geprinte behuizing. Op een eigen PCB bevinden zich een microcontroller, GPS-RTK-module, barometer, IMU-sensor en LoRa-zender.",
    "De GPS-RTK-module bepaalt de positie, de barometer meet de hoogte en de IMU meet beweging en oriëntatie. De microcontroller verzamelt deze gegevens en verstuurt ze via LoRa naar een laptop. Op de laptop worden de gegevens opgeslagen en weergegeven in een eenvoudige 3D-omgeving.",
    "De behuizing krijgt een dempingsmechanisme, bijvoorbeeld met veren of flexibel materiaal. Hiermee onderzoek ik of schokken en trillingen minder invloed hebben op de metingen.",
    "Als uitbreiding maak ik een stilstaande vliegtuig-mock-up. De laptop berekent op basis van de ontvangen oriëntatiegegevens een eenvoudige correctie. Deze wordt naar een tweede microcontroller gestuurd, die twee servo’s beweegt. De servo’s beelden het rolroer en hoogteroer van een vliegtuig uit. Zo toon ik dat sensorgegevens niet alleen zichtbaar gemaakt, maar ook gebruikt kunnen worden voor een eenvoudige aansturing.",
]
for text in paragraphs:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)

heading("Blokschema")
intro = doc.add_paragraph("Meetmodule en verwerking/aansturing")
intro.paragraph_format.space_after = Pt(2)
intro.runs[0].italic = True

table = doc.add_table(rows=3, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [3.0, 0.65, 3.45, 0.65, 4.15]
for row in table.rows:
    for i, cell in enumerate(row.cells):
        cell.width = Cm(widths[i])

add_box(table.cell(0, 0), "GPS-RTK\nBarometer\nIMU", "EAF2F8")
table.cell(0, 1).text = "→"
table.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
table.cell(0, 1).paragraphs[0].runs[0].font.size = Pt(16)
add_box(table.cell(0, 2), "Microcontroller\n+ LoRa-zender\n(in behuizing met demping)", "DDEBF7")
table.cell(0, 3).text = "→"
table.cell(0, 3).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
table.cell(0, 3).paragraphs[0].runs[0].font.size = Pt(16)
add_box(table.cell(0, 4), "LoRa-ontvanger → Laptop\nOpslag · 3D-weergave\nBerekening correctie", "E2F0D9")

for cell in table.rows[1].cells:
    cell.text = ""
table.cell(1, 4).text = "↓"
table.cell(1, 4).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
table.cell(1, 4).paragraphs[0].runs[0].font.size = Pt(15)

for cell in table.rows[2].cells:
    cell.text = ""
add_box(table.cell(2, 4), "Tweede microcontroller → 2 servo’s\nRolroer + hoogteroer\nVliegtuig-mock-up", "FCE4D6")

heading("Doel")
p = doc.add_paragraph("Een werkend prototype maken dat positie, hoogte en oriëntatie meet, de gegevens draadloos verstuurt, ze in 3D weergeeft en via een mock-up demonstreert hoe zulke gegevens gebruikt kunnen worden voor aansturing.")
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

doc.save(OUTPUT)
print(OUTPUT)
