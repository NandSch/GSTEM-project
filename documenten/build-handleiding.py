# -*- coding: utf-8 -*-
"""
Bouwt het Word-document van de handleiding uit de markdown-bron.

Gebruik:  python documenten/build-handleiding.py
Bron:     documenten/Handleiding-meettoestel.md
Doel:     documenten/Handleiding-meettoestel.docx
"""

import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

HIER = os.path.dirname(os.path.abspath(__file__))
BRON = os.path.join(HIER, "Handleiding-meettoestel.md")
DOEL = os.path.join(HIER, "Handleiding-meettoestel.docx")

CALLOUT_KLEUR = {
    "screenshot": "FCE8E6",   # zacht rood -> actie voor jou
    "tip": "E8F0FE",          # zacht blauw
    "warning": "FFF4E5",      # zacht oranje
    "info": "EAF3EA",         # zacht groen
}

CALLOUT_LABEL = {
    "screenshot": "In te voegen",
    "tip": "Tip",
    "warning": "Let op",
    "info": "Info",
}


def shade_cell(cell, hex_kleur):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_kleur)
    tc_pr.append(shd)


def add_runs(paragraph, tekst):
    """Zet **vet**, *cursief* en ***vet cursief*** om naar echte opmaak."""
    for deel in re.split(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*[^*]+?\*)", tekst):
        if not deel:
            continue
        vet = cursief = False
        if deel.startswith("***") and deel.endswith("***"):
            inhoud, vet, cursief = deel[3:-3], True, True
        elif deel.startswith("**") and deel.endswith("**"):
            inhoud, vet = deel[2:-2], True
        elif deel.startswith("*") and deel.endswith("*"):
            inhoud, cursief = deel[1:-1], True
        else:
            inhoud = deel
        run = paragraph.add_run(inhoud)
        run.bold = vet
        run.italic = cursief


def add_callout(doc, soort, titel, body_regels):
    kleur = CALLOUT_KLEUR.get(soort, "F1F1F1")
    label = CALLOUT_LABEL.get(soort, soort.capitalize())
    tabel = doc.add_table(rows=1, cols=1)
    tabel.alignment = WD_TABLE_ALIGNMENT.CENTER
    cel = tabel.cell(0, 0)
    shade_cell(cel, kleur)

    eerste = cel.paragraphs[0]
    run = eerste.add_run(f"[{label}] {titel}")
    run.bold = True
    for regel in body_regels:
        p = cel.add_paragraph()
        add_runs(p, regel)

    doc.add_paragraph()


def add_table(doc, regels):
    rijen = []
    for regel in regels:
        cellen = [c.strip() for c in regel.strip().strip("|").split("|")]
        rijen.append(cellen)
    if len(rijen) >= 2 and set(rijen[1][0]) <= set("-: "):
        kop, body = rijen[0], rijen[2:]
    else:
        kop, body = None, rijen

    kolommen = max(len(r) for r in rijen)
    tabel = doc.add_table(rows=0, cols=kolommen)
    tabel.style = "Light Grid Accent 1"
    if kop:
        row = tabel.add_row().cells
        for i in range(kolommen):
            tekst = kop[i] if i < len(kop) else ""
            row[i].text = ""
            add_runs(row[i].paragraphs[0], "**%s**" % tekst if tekst else "")
    for rij in body:
        row = tabel.add_row().cells
        for i in range(kolommen):
            tekst = rij[i] if i < len(rij) else ""
            row[i].text = ""
            add_runs(row[i].paragraphs[0], tekst)
    doc.add_paragraph()


def bouw():
    with open(BRON, encoding="utf-8") as fh:
        regels = fh.read().splitlines()

    doc = Document()

    # Opmaak afgestemd op het referentiedocument GStem-Specificaties:
    # body Calibri 12, secties vet-cursief 17, tussenkoppen vet-cursief 13.
    stijl = doc.styles["Normal"]
    stijl.font.name = "Calibri"
    stijl.font.size = Pt(12)

    kop1 = doc.styles["Heading 1"]
    kop1.font.name = "Calibri"
    kop1.font.size = Pt(17)
    kop1.font.bold = True
    kop1.font.italic = True
    kop1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    tussenkop = doc.styles["Subtitle"]
    tussenkop.font.name = "Calibri"
    tussenkop.font.size = Pt(13)
    tussenkop.font.bold = True
    tussenkop.font.italic = True
    tussenkop.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    titelstijl = doc.styles["Title"]
    titelstijl.font.name = "Calibri"

    i = 0
    n = len(regels)
    while i < n:
        regel = regels[i]
        kaal = regel.strip()

        if not kaal:
            i += 1
            continue

        # Kopjes
        if kaal.startswith("### "):
            doc.add_paragraph(kaal[4:].strip(), style="Subtitle")
            i += 1
            continue
        if kaal.startswith("## "):
            doc.add_paragraph(kaal[3:].strip(), style="Heading 1")
            i += 1
            continue
        # Titel: gewone regel met de projectnaam vet-cursief (zoals de referentie)
        if kaal.startswith("# "):
            par = doc.add_paragraph()
            add_runs(par, kaal[2:].strip())
            for r in par.runs:
                r.font.size = Pt(13)
            i += 1
            continue

        # Afbeelding met optioneel onderschrift
        m = re.match(r"!\[(.*?)\]\((.+?)\)\s*$", kaal)
        if m:
            onderschrift, pad = m.group(1).strip(), m.group(2).strip()
            volledig = pad if os.path.isabs(pad) else os.path.join(HIER, pad)
            doc.add_picture(volledig, width=Inches(6.0))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if onderschrift:
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cap.add_run(onderschrift)
                run.italic = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            i += 1
            continue

        # Horizontale lijn
        if set(kaal) <= set("-*_") and len(kaal) >= 3:
            i += 1
            continue

        # Callout-blok
        if kaal.startswith("> "):
            blok = []
            while i < n and regels[i].strip().startswith(">"):
                blok.append(regels[i].strip().lstrip(">").strip())
                i += 1
            eerste = blok[0]
            m = re.match(r"\[!(\w+)\]\s*(.*)", eerste)
            if m:
                soort, titel = m.group(1).lower(), m.group(2)
                body = [b for b in blok[1:] if b]
            else:
                soort, titel, body = "info", eerste, [b for b in blok[1:] if b]
            add_callout(doc, soort, titel, body)
            continue

        # Tabel
        if kaal.startswith("|"):
            blok = []
            while i < n and regels[i].strip().startswith("|"):
                blok.append(regels[i].strip())
                i += 1
            add_table(doc, blok)
            continue

        # Opsomming
        if kaal.startswith("- "):
            inhoud = kaal[2:].strip()
            if inhoud.startswith("[ ] "):
                par = doc.add_paragraph(style="List Bullet")
                par.add_run("\u2610 ")
                add_runs(par, inhoud[4:])
            else:
                par = doc.add_paragraph(style="List Bullet")
                add_runs(par, inhoud)
            i += 1
            continue

        # Cursieve slotregel
        if kaal.startswith("*") and kaal.endswith("*") and not kaal.startswith("**"):
            par = doc.add_paragraph()
            run = par.add_run(kaal.strip("*"))
            run.italic = True
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            i += 1
            continue

        # Gewone paragraaf
        par = doc.add_paragraph()
        add_runs(par, kaal)
        i += 1

    doc.save(DOEL)
    print("Geschreven:", DOEL)


if __name__ == "__main__":
    bouw()
