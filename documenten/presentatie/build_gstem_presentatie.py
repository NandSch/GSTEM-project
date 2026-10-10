
# -*- coding: utf-8 -*-
"""Bouwt Presentatie-GSTEM.pptx (16:9) voor het G-STEM-P meettoestelproject."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw
import os

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
AFB = os.path.join(ROOT, "specificaties", "afbeeldingen")

INK   = RGBColor(0x1B, 0x1F, 0x3B)
INK2  = RGBColor(0x33, 0x38, 0x55)
ACC   = RGBColor(0x0E, 0x6B, 0x5C)
ACC_L = RGBColor(0xE7, 0xF3, 0xF0)
ACC_L2= RGBColor(0xF1, 0xF8, 0xF6)
GRIJS = RGBColor(0x5A, 0x60, 0x72)
LIJN  = RGBColor(0xB9, 0xC2, 0xC9)
WIT   = RGBColor(0xFF, 0xFF, 0xFF)
ROOD  = RGBColor(0xC2, 0x40, 0x34)
ROOD_L= RGBColor(0xFB, 0xEC, 0xEA)
GEEL_L= RGBColor(0xFD, 0xF3, 0xE0)
GEEL  = RGBColor(0xB0, 0x7D, 0x10)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height

def add_slide():
    return prs.slides.add_slide(BLANK)

def txt(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=6, wrap=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        for (t, size, bold, color, italic) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; r.font.italic = italic
            r.font.name = "Calibri"
    return tb

def R(t, size=18, bold=False, color=INK2, italic=False):
    return (t, size, bold, color, italic)

def box(slide, x, y, w, h, fill=ACC_L, line=None, radius=None, shadow=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shape_type, x, y, w, h)
    if radius is not None:
        try: sp.adjustments[0] = radius
        except Exception: pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = shadow
    return sp

def header(slide, kicker, title):
    box(slide, 0, 0, SW, Inches(1.18), fill=WIT)
    box(slide, Inches(0.55), Inches(0.32), Inches(0.14), Inches(0.62), fill=ACC)
    txt(slide, Inches(0.9), Inches(0.22), Inches(10.5), Inches(0.3), [[R(kicker.upper(), 11, True, ACC)]])
    txt(slide, Inches(0.9), Inches(0.5), Inches(11.5), Inches(0.6), [[R(title, 28, True, INK)]])
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.55), Inches(1.18), SW - Inches(0.55), Inches(1.18))
    ln.line.color.rgb = LIJN; ln.line.width = Pt(1)

def chip(slide, x, y, w, h, label_, sub=None, fill=ACC_L, fg=INK, big=False, line=None, tsize=None, ssize=None):
    box(slide, x, y, w, h, fill=fill, radius=0.14, line=line)
    if sub:
        txt(slide, x + Inches(0.15), y + Inches(0.08), w - Inches(0.3), h - Inches(0.16),
            [[R(label_, tsize or (15 if big else 13), True, fg)], [R(sub, ssize or (11.5 if big else 10.5), False, fg)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=2)
    else:
        txt(slide, x + Inches(0.1), y, w - Inches(0.2), h,
            [[R(label_, tsize or (14 if big else 12.5), True, fg)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def arrow(slide, x1, y1, x2, y2, color=ACC, width=2.25, dashed=False):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    ln.line.color.rgb = color; ln.line.width = Pt(width)
    lnEl = ln.line._get_or_add_ln()
    if dashed:
        d = lnEl.makeelement(qn('a:prstDash'), {'val': 'dash'}); lnEl.append(d)
    lnEl.append(lnEl.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    return ln

def label(slide, x, y, w, text, size=11, color=GRIJS, align=PP_ALIGN.CENTER):
    txt(slide, x, y, w, Inches(0.3), [[R(text, size, False, color, True)]], align=align)

def picture_fit(slide, path, x, y, max_w, max_h, frame=True):
    im = Image.open(path); iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    px = x + int((max_w - w) / 2); py = y + int((max_h - h) / 2)
    if frame:
        box(slide, px - Emu(19050), py - Emu(19050), w + Emu(38100), h + Emu(38100), fill=WIT, line=LIJN, radius=0.05)
    return slide.shapes.add_picture(path, px, py, w, h)

def make_device_image(path):
    W, H = 980, 640
    im = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(im)
    d.ellipse([90, 545, 890, 615], fill=(235, 238, 241))
    d.rounded_rectangle([120, 210, 860, 520], radius=38, fill=(31, 47, 63))
    d.rounded_rectangle([120, 170, 860, 480], radius=38, fill=(41, 61, 82))
    d.rounded_rectangle([150, 195, 830, 455], radius=26, fill=(52, 76, 101))
    d.ellipse([172, 222, 196, 246], fill=(80, 220, 150))
    d.rounded_rectangle([214, 224, 420, 244], radius=9, fill=(30, 44, 60))
    d.text((230, 227), "G-STEM  MEETMODULE", fill=(200, 214, 226))
    d.rectangle([700, 118, 716, 178], fill=(58, 76, 95))
    d.rounded_rectangle([668, 84, 748, 126], radius=14, fill=(222, 226, 231), outline=(160, 168, 178), width=2)
    for yy in (96, 106, 116):
        d.line([678, yy, 738, yy], fill=(180, 186, 195), width=3)
    d.rounded_rectangle([180, 300, 320, 400], radius=16, fill=(196, 200, 206), outline=(150, 156, 164), width=2)
    d.ellipse([214, 330, 286, 370], fill=(226, 229, 234), outline=(150, 156, 164), width=2)
    d.text((210, 408), "GNSS-antenne", fill=(214, 224, 233))
    for i in range(4):
        d.rounded_rectangle([360 + i * 46, 300, 388 + i * 46, 396], radius=8, fill=(37, 55, 74))
    d.rounded_rectangle([560, 300, 810, 400], radius=14, fill=(66, 92, 118))
    for i, lab in enumerate(["TX", "RX", "GND"]):
        cx = 600 + i * 74
        d.rounded_rectangle([cx, 330, cx + 44, 370], radius=8, fill=(222, 226, 231), outline=(140, 148, 158), width=2)
        d.text((cx + 12, 375), lab, fill=(210, 220, 230))
    d.text((610, 408), "Aansluiting voertuig", fill=(214, 224, 233))
    d.rounded_rectangle([118, 470, 158, 500], radius=6, fill=(18, 28, 40))
    d.ellipse([126, 478, 150, 492], fill=(90, 96, 106))
    im.save(path)

DEVICE_IMG = os.path.join(ROOT, "presentatie", "device-illustratie.png")
make_device_image(DEVICE_IMG)

# ---------- SLIDE 1 - titel ----------
s = add_slide()
box(s, 0, 0, SW, SH, fill=WIT)
box(s, 0, 0, SW, Inches(0.16), fill=ACC)
txt(s, Inches(0.9), Inches(2.6), Inches(6.4), Inches(2.4), [
    [R("Positie- en beweging", 40, True, INK)],
    [R("meettoestel met LoRa", 40, True, INK)],
    [R("G-STEM-P  2026-2027", 16, True, ACC)],
], space_after=10)
txt(s, Inches(0.9), Inches(5.6), Inches(6.0), Inches(0.6), [[R("Nand Schoovaerts", 20, True, INK2)]])
picture_fit(s, DEVICE_IMG, Inches(7.5), Inches(1.7), Inches(5.2), Inches(3.9))
label(s, Inches(7.5), Inches(5.75), Inches(5.2), "Voorgesteld uitzicht van het meettoestel (concept)", 11)

# ---------- SLIDE 2 - inleiding ----------
s = add_slide()
header(s, "Inleiding", "Wat is het project?")
txt(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(1.3), [
    [R("Een klein meettoestel op een bewegend voertuig meet vier grootheden en stuurt ze via LoRa naar de laptop-app.", 17, False, INK2)],
], space_after=0)
labels = [
    ("Richting", "hoe schuin of recht het toestel staat (graden)"),
    ("Snelheid", "hoe snel het beweegt (km/u)"),
    ("Hoogte", "nauwkeurig tot op 1,5 meter"),
    ("Locatie", "nauwkeurig tot op 0,5 meter"),
]
cw = Inches(2.75); gapw = Inches(0.27); x0 = Inches(0.9)
for i, (t, sub) in enumerate(labels):
    x = x0 + i * (cw + gapw)
    chip(s, x, Inches(3.1), cw, Inches(1.15), t, sub, big=True)
txt(s, Inches(0.9), Inches(4.55), Inches(11.5), Inches(0.4),
    [[R("Draadloze ontvanger in USB-stickvorm. Bereik: tot 4 km.", 15, False, INK2)]])
chain = ["Meettoestel", "LoRa", "USB-ontvanger", "Laptop-app"]
bw = Inches(2.3); bh = Inches(0.62)
for i, t in enumerate(chain):
    x = Inches(0.9) + i * Inches(3.05)
    chip(s, x, Inches(5.6), bw, bh, t, fill=WIT, fg=INK, line=LIJN, tsize=13.5)
    if i < len(chain) - 1:
        arrow(s, x + bw + Inches(0.08), Inches(5.6) + bh / 2, Inches(0.9) + (i + 1) * Inches(3.05) - Inches(0.08), Inches(5.6) + bh / 2, color=ACC, width=2)
label(s, Inches(0.9), Inches(6.5), Inches(11.5), "De volledige keten volgt later.", 11)

# ---------- SLIDE 3 - sensoren ----------
s = add_slide()
header(s, "Hardware", "Welke sensoren gebruik ik?")
sens = [
    ("IMU (9-DoF)", "Richting, versnelling en hoeksnelheid."),
    ("Barometer", "Luchtdruk en hoogte."),
    ("RTK-GNSS", "Locatie tot op 0,5 m, met eigen antenne."),
    ("ESP32-S3 + LoRa", "Leest alles uit en zendt draadloos."),
]
cw = Inches(2.75); gapw = Inches(0.27); x0 = Inches(0.9); y0 = Inches(1.6)
for i, (t, sub) in enumerate(sens):
    x = x0 + i * (cw + gapw)
    box(s, x, y0, cw, Inches(2.5), fill=ACC_L2, radius=0.1)
    txt(s, x + Inches(0.18), y0 + Inches(0.22), cw - Inches(0.36), Inches(2.1),
        [[R(t, 16, True, INK)], [R(sub, 12.5, False, INK2)]], space_after=8)
y1 = Inches(4.5)
chip(s, Inches(0.9), y1, Inches(5.6), Inches(0.95), "Voeding: 7,4 V-accu en LED-status",
     sub="start automatisch met het voertuig", big=True)
chip(s, Inches(6.8), y1, Inches(5.6), Inches(0.95), "3D-geprinte behuizing met demping",
     sub="beschermt tegen schokken en trillingen", big=True)
label(s, Inches(0.9), Inches(5.85), Inches(11.5),
      "Een ESP32-S3 leest alle sensoren uit.", 12)

# ---------- SLIDE 4 - app deel 1 ----------
s = add_slide()
header(s, "De applicatie - deel 1", "Verbinding controleren en live data zien")
txt(s, Inches(0.9), Inches(1.45), Inches(11.5), Inches(0.5),
    [[R("De app start vanzelf wanneer de USB-ontvanger wordt ingestoken. Op elk scherm staat de status van de ontvanger en het meettoestel.", 14.5, False, INK2)]])
picture_fit(s, os.path.join(AFB, "image2-scherm1-setup.png"), Inches(0.9), Inches(2.1), Inches(5.4), Inches(2.1))
txt(s, Inches(0.9), Inches(4.35), Inches(5.4), Inches(1.6), [
    [R("Scherm 1 - Verbindingscontrole", 15, True, INK)],
    [R("Controleert de verbindingen. De gebruiker klikt op OK om door te gaan.", 12.5, False, INK2)],
], space_after=6)
picture_fit(s, os.path.join(AFB, "image3-scherm2-kaart.png"), Inches(6.9), Inches(2.1), Inches(5.5), Inches(2.1))
txt(s, Inches(6.9), Inches(4.35), Inches(5.5), Inches(1.6), [
    [R("Scherm 2 - Kaart en live data", 15, True, INK)],
    [R("3D-kaart met satellietfotos toont de afgelegde weg en kijkrichting; een tabel toont alle losse meetwaarden.", 12.5, False, INK2)],
], space_after=6)

# ---------- SLIDE 5 - app deel 2 ----------
s = add_slide()
header(s, "De applicatie - deel 2", "Besturen vanuit de app")
txt(s, Inches(0.9), Inches(1.45), Inches(11.5), Inches(0.5),
    [[R("Via de knop Besturing kies je een voertuigspecifieke bedieningspagina; elk voertuig heeft zijn eigen bediening.", 14.5, False, INK2)]])
picture_fit(s, os.path.join(AFB, "image4-scherm3-besturing.png"), Inches(0.9), Inches(2.1), Inches(5.4), Inches(2.1))
txt(s, Inches(0.9), Inches(4.35), Inches(5.4), Inches(1.6), [
    [R("Scherm 3 - Besturingsmenu", 15, True, INK)],
    [R("Drie applicatie-iconen, een per voertuigtype. Elk icoon opent de bijbehorende besturingspagina.", 12.5, False, INK2)],
], space_after=6)
picture_fit(s, os.path.join(AFB, "image5-scherm4-auto.png"), Inches(6.9), Inches(2.1), Inches(5.5), Inches(2.1))
txt(s, Inches(6.9), Inches(4.35), Inches(5.5), Inches(1.6), [
    [R("Scherm 4 - Besturing (auto)", 15, True, INK)],
    [R("Stuur door te klikken en slepen; het gaspedaal indrukken laat het voertuig vooruit rijden.", 12.5, False, INK2)],
], space_after=6)

# ---------- SLIDE 6 - communicatie (verbeterd diagram) ----------
s = add_slide()
header(s, "Communicatie", "Hoe praten alle onderdelen met elkaar?")
# sensoren links
sx, sw, sh_ = Inches(0.9), Inches(2.5), Inches(0.62)
sens_y = [Inches(1.75), Inches(2.55), Inches(3.35)]
for i, t in enumerate(["IMU (9-DoF)", "Barometer", "RTK-GNSS"]):
    chip(s, sx, sens_y[i], sw, sh_, t, big=True, tsize=14)
# ESP32 midden-links
ex, ew, eh = Inches(4.15), Inches(2.55), Inches(1.5)
chip(s, ex, Inches(2.15), ew, eh, "ESP32-S3 + LoRa", sub="leest alle sensoren uit en bundelt de meetdata", big=True, tsize=14.5, ssize=11)
esp_cy = Inches(2.15) + eh / 2
# pijlen sensoren -> ESP
for i in range(3):
    arrow(s, sx + sw + Inches(0.06), sens_y[i] + sh_ / 2, ex - Inches(0.06), esp_cy, color=ACC, width=1.75)
label(s, Inches(3.32), Inches(2.18), Inches(0.9), "I2C", 10.5, align=PP_ALIGN.CENTER)
label(s, Inches(3.32), Inches(3.62), Inches(0.9), "UART", 10.5, align=PP_ALIGN.CENTER)
# USB-ontvanger en laptop rechts
ux, uw, uh = Inches(7.75), Inches(2.15), Inches(0.7)
chip(s, ux, Inches(2.55), uw, uh, "USB-ontvanger", sub="tweede LoRa-module", big=True, tsize=13.5, ssize=10.5)
lx, lw = Inches(10.35), Inches(2.15)
chip(s, lx, Inches(2.55), lw, uh, "Laptop-app", sub="kaart + besturing", big=True, tsize=13.5, ssize=10.5)
# LoRa dubbele pijl
lora_y = Inches(2.9)
arrow(s, ex + ew + Inches(0.06), lora_y - Inches(0.14), ux - Inches(0.06), lora_y - Inches(0.14), color=ROOD, width=2.5)
arrow(s, ux - Inches(0.06), lora_y + Inches(0.14), ex + ew + Inches(0.06), lora_y + Inches(0.14), color=ROOD, width=2.5)
label(s, Inches(6.62), Inches(2.1), Inches(1.2), "LoRa", 12, color=ROOD)
label(s, Inches(6.72), Inches(3.3), Inches(1.05), "tot 4 km", 10.5, color=ROOD)
# USB pijl
arrow(s, ux + uw + Inches(0.06), Inches(2.9), lx - Inches(0.06), Inches(2.9), color=ACC, width=2)
label(s, Inches(9.82), Inches(2.2), Inches(0.62), "USB", 10.5)
# commandos terug (gestippeld) laptop -> USB -> ESP
arrow(s, lx - Inches(0.06), Inches(3.62), ex + ew + Inches(0.06), Inches(3.62), color=GRIJS, width=1.5, dashed=True)
label(s, Inches(6.35), Inches(3.72), Inches(4.0), "stuurcommando's terug", 10.5, color=GRIJS)
# Arduino onder
ax, aw, ah = Inches(4.15), Inches(2.55), Inches(0.85)
chip(s, ax, Inches(4.75), aw, ah, "Arduino (voertuig)", sub="stuurt motoren en servo's aan", big=True, tsize=13.5, ssize=11)
arrow(s, ex + Inches(0.75), Inches(3.65) + Inches(0.15), ax + Inches(0.75), Inches(4.75) - Inches(0.06), color=ACC, width=2)
label(s, Inches(5.0), Inches(4.12), Inches(1.7), "UART TX/RX", 10.5, align=PP_ALIGN.LEFT)
arrow(s, ax + Inches(1.55), Inches(4.75) - Inches(0.06), ex + Inches(1.55), Inches(3.8), color=GRIJS, width=1.5, dashed=True)
# uitlegblok onder
box(s, Inches(0.9), Inches(5.95), Inches(11.53), Inches(0.85), fill=ACC_L2, radius=0.12)
txt(s, Inches(1.15), Inches(6.08), Inches(11.0), Inches(0.6),
    [[R("Meetdata stromen van de sensoren naar de app; stuurcommando's gaan via dezelfde weg terug.", 13, True, INK)],
     [R("Bij verbindingsverlies of ongeldige commando's stopt het voertuig veilig.", 12, False, INK2)]], space_after=3)

# ---------- SLIDE 7 - planning en deadlines ----------
s = add_slide()
header(s, "Planning", "Deadlines en werffases")
rows = [
    ("13/10/2026",       "Voorlopige presentatie SVL (5 min)",              False),
    ("nov-dec 2026",     "Onderdelen bestellen en apart testen",            True),
    ("01/12/2026",       "Evaluatie SVZ met mentor",                        False),
    ("dec 2026-feb 2027","Demo-app: kaart en live data",                    True),
    ("feb-mei 2027",     "Volledige app: besturing erbij",                  True),
    ("19-23/03/2027",    "SVZ-presentatie (10 min)",                        False),
    ("mrt-mei 2027",     "Toestel monteren, testen en kalibreren",          True),
    ("22/05/2027",       "Opendeurdag: demo + scriptie",                    False),
    ("21/06/2027",       "Juryverdediging",                                 False),
]
ry = Inches(1.5)
rh = Inches(0.5); rgap = Inches(0.09)
for (datum, what, werk) in rows:
    if werk:
        chip(s, Inches(0.9), ry, Inches(2.6), rh, datum, fill=ACC, fg=WIT, tsize=13)
    else:
        chip(s, Inches(0.9), ry, Inches(2.6), rh, datum, fill=ACC_L, fg=INK, tsize=13)
    txt(s, Inches(3.75), ry, Inches(8.6), rh, [[R(what, 14.5, werk, INK if werk else INK2)]],
        anchor=MSO_ANCHOR.MIDDLE)
    ry = ry + rh + rgap
label(s, Inches(0.9), Inches(6.75), Inches(11.5),
      "Groene balk = werffase   |   lichte balk = vaste mijlpaal uit de projectplanning", 11, align=PP_ALIGN.LEFT)

# ---------- SLIDE 8 - bestellijst ----------
s = add_slide()
header(s, "Budget", "Bestellijst per groep")
groups = [
    ("Sensoren (IMU, barometer, RTK-GNSS)", 79.92, "59%"),
    ("Rekenkern + LoRa (2x XIAO-kit)",      30.26, "22%"),
    ("Bedrading en losse elektronica",      19.56, "15%"),
    ("Antennes en RF (pigtail)",             3.57, "3%"),
    ("Voeding (LDO, zekering, TVS)",         1.27, "1%"),
]
maxw = Inches(5.6)
gy = Inches(1.75)
for (name, bedrag, pct) in groups:
    txt(s, Inches(0.9), gy, Inches(3.9), Inches(0.55), [[R(name, 13, True, INK)]], anchor=MSO_ANCHOR.MIDDLE)
    w = max(int(maxw * bedrag / 79.92), Inches(0.22))
    box(s, Inches(4.95), gy + Inches(0.08), w, Inches(0.4), fill=ACC, radius=0.18)
    txt(s, Inches(4.95) + w + Inches(0.12), gy, Inches(1.6), Inches(0.55),
        [[R(("{:.2f}".format(bedrag)).replace(",", "X").replace(".", ",").replace("X", ".") + " \u20ac  " + pct, 12.5, True, INK2)]],
        anchor=MSO_ANCHOR.MIDDLE)
    gy = gy + Inches(0.88)
box(s, Inches(0.9), Inches(6.05), Inches(11.53), Inches(1.0), fill=ACC_L2, radius=0.1)
txt(s, Inches(1.15), Inches(6.16), Inches(5.4), Inches(0.8),
    [[R("Totaal: \u20ac 134,58", 17, True, ACC)],
     [R("+ \u20ac 15,70 losse GNSS-antenne (alleen als de kit er geen meelevert) \u2192 \u20ac 150,28", 11.5, False, INK2)]], space_after=3)
txt(s, Inches(6.9), Inches(6.16), Inches(5.4), Inches(0.8),
    [[R("Grootste kostenpost: de sensoren", 13.5, True, INK)],
     [R("vooral RTK-GNSS (\u20ac 36,99) en IMU (\u20ac 32,05) samen \u20ac 69,04. Accu, Arduino, servo's en kabels heeft de gebruiker al.", 11.5, False, INK2)]], space_after=3)

# ---------- SLIDE 9 - slot ----------
s = add_slide()
box(s, 0, 0, SW, SH, fill=WIT)
box(s, 0, SH - Inches(0.16), SW, Inches(0.16), fill=ACC)
txt(s, Inches(0.9), Inches(2.7), Inches(11.5), Inches(1.6), [
    [R("Bedankt voor het luisteren.", 34, True, INK)],
    [R("Vragen zijn welkom.", 20, False, INK2)],
], space_after=10)
txt(s, Inches(0.9), Inches(4.6), Inches(11.5), Inches(0.5),
    [[R("Nand Schoovaerts  G-STEM-P  2026-2027", 15, False, GRIJS)]])

out = os.path.join(ROOT, "presentatie", "Presentatie-GSTEM.pptx")
prs.save(out)
print("OK", out)
